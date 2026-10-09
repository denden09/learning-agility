import base64
import io
import math
import os
import posixpath
import re
import shutil
import textwrap
import zipfile
from datetime import date, datetime
from html import escape as html_escape
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st
import xml.etree.ElementTree as ET
from openpyxl import load_workbook


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Learning System",
    page_icon="📙",
    layout="wide",
)


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

# NOTE: Linux (Streamlit Cloud) is CASE-SENSITIVE.
# Nama file / folder harus persis sama dengan yang ada di repo.
MASTER_FILE = (
    BASE_DIR
    / "data"
    / "master"
    / "Learning agility self assesment & superior_01.10.2026 1.xlsx"
)

OUTPUT_DIR = BASE_DIR / "generated"

BACKGROUND_IMAGE = BASE_DIR / "public" / "abstract-orange-bg.png"


def get_background_base64(image_path):

    if not image_path.exists():
        return None

    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode("utf-8")


BACKGROUND_BASE64 = get_background_base64(BACKGROUND_IMAGE)


# ============================================================
# DIMENSIONS
# ============================================================

DIMENSIONS = {
    "mental_agility": "mental agility",
    "people_agility": "people agility",
    "change_agility": "change agility",
    "result_agility": "result agility",
    "self_awareness": "self awareness",
}


# ============================================================
# IDENTITY
# ============================================================

IDENTITY_INPUT_CELLS = [
    "D2", "D3", "D4", "D5", "D6", "D7", "D8", "D9", "D10", "D12",
]


# ============================================================
# ASSESSMENT COLUMNS
# ============================================================

ASSESSMENT_INPUT_COLUMNS = {
    "self_rating": 9,
    "self_comment": 10,
    "superior_rating": 11,
    "superior_comment": 12,
}

# Row 13 is the first assessment row in every dimension sheet.
# The number of questions is NOT fixed: the Master workbook is the source
# of truth, so each sheet may have a different number of indicators.
ASSESSMENT_START_ROW = 13

ASSESSMENT_SHEETS = list(DIMENSIONS.values())


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(value):

    if pd.isna(value):
        return ""

    value = str(value)
    value = value.replace("\xa0", " ")
    value = value.replace("\n", " ")
    value = value.replace("\r", " ")
    value = re.sub(r"\s+", " ", value)

    return value.strip().lower()


# ============================================================
# CLEAN VALUE
# ============================================================

def clean_value(value):

    if pd.isna(value):
        return ""

    return str(value).strip()


# ============================================================
# CLEAN EXCEL VALUE
# ============================================================

def clean_excel_value(value):

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return None

    return value


# ============================================================
# RENDER HTML
# ============================================================

def render_html(html):

    clean_html = " ".join(
        line.strip()
        for line in html.splitlines()
        if line.strip()
    )

    st.markdown(clean_html, unsafe_allow_html=True)


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(df, candidates):

    if df is None or df.empty:
        return None

    normalized_columns = {
        normalize_text(col): col
        for col in df.columns
    }

    for candidate in candidates:

        candidate_normalized = normalize_text(candidate)

        if candidate_normalized in normalized_columns:
            return normalized_columns[candidate_normalized]

    return None


# ============================================================
# UPLOAD STATE HELPERS
#
# File disimpan sebagai bytes di session_state.
# ============================================================

def store_uploaded_file(role, uploaded_file):

    st.session_state[f"{role}_file_bytes"] = uploaded_file.getvalue()
    st.session_state[f"{role}_file_name"] = uploaded_file.name


def get_stored_file(role):

    file_bytes = st.session_state.get(f"{role}_file_bytes")
    file_name = st.session_state.get(f"{role}_file_name")

    if file_bytes is None or file_name is None:
        return None

    file_buffer = io.BytesIO(file_bytes)
    file_buffer.name = file_name

    return file_buffer


def clear_uploaded_file(role):

    st.session_state[f"{role}_file_bytes"] = None
    st.session_state[f"{role}_file_name"] = None


# ============================================================
# READ MASTER STRUCTURE
# ============================================================

def read_master_indicators(master_file):

    if not master_file.exists():

        raise FileNotFoundError(
            "Master Template tidak ditemukan.\n\n"
            f"Path:\n{master_file}\n\n"
            "Pastikan file berada di folder:\n"
            "data/master/"
        )

    master_wb = load_workbook(master_file, data_only=False)

    indicator_master = []

    try:

        for dimension_key, sheet_name in DIMENSIONS.items():

            if sheet_name not in master_wb.sheetnames:

                raise ValueError(
                    f"Sheet '{sheet_name}' "
                    "tidak ditemukan di Master."
                )

            ws = master_wb[sheet_name]

            # Read every populated indicator row from the Master.
            # Do not hardcode 10 questions per dimension.
            for row in range(ASSESSMENT_START_ROW, ws.max_row + 1):

                indicator = ws.cell(row, 2).value
                question = ws.cell(row, 3).value

                # A completely empty row is not an indicator.
                if clean_value(indicator) == "" and clean_value(question) == "":
                    continue

                if clean_value(question) == "":
                    raise ValueError(
                        f"Sheet '{sheet_name}' row {row}: "
                        "Indicator terisi tetapi Question kosong."
                    )

                indicator_no = row - ASSESSMENT_START_ROW + 1

                levels = {
                    1: ws.cell(row, 4).value,
                    2: ws.cell(row, 5).value,
                    3: ws.cell(row, 6).value,
                    4: ws.cell(row, 7).value,
                    5: ws.cell(row, 8).value,
                }

                indicator_master.append({
                    "dimension": dimension_key,
                    "sheet_name": sheet_name,
                    "indicator_no": indicator_no,
                    "master_row": row,
                    "indicator_id": f"{dimension_key}_{indicator_no:02d}",
                    "indicator": indicator,
                    "question": question,
                    "question_normalized": normalize_text(question),
                    "level_1": levels[1],
                    "level_2": levels[2],
                    "level_3": levels[3],
                    "level_4": levels[4],
                    "level_5": levels[5],
                    "level_1_normalized": normalize_text(levels[1]),
                    "level_2_normalized": normalize_text(levels[2]),
                    "level_3_normalized": normalize_text(levels[3]),
                    "level_4_normalized": normalize_text(levels[4]),
                    "level_5_normalized": normalize_text(levels[5]),
                })

    finally:

        master_wb.close()

    indicator_master_df = pd.DataFrame(indicator_master)

    expected_question_count = len(indicator_master_df)

    if expected_question_count == 0:
        raise ValueError(
            "Tidak ada indikator yang ditemukan pada Master. "
            f"Pastikan data dimulai dari row {ASSESSMENT_START_ROW}."
        )

    if indicator_master_df["indicator_id"].nunique() != expected_question_count:

        raise ValueError(
            "Terdapat duplicate indicator_id pada Master."
        )

    if indicator_master_df["question_normalized"].nunique() != expected_question_count:

        raise ValueError(
            "Terdapat duplicate question pada Master."
        )

    for col in ["level_1", "level_2", "level_3", "level_4", "level_5"]:

        if indicator_master_df[col].isna().sum() != 0:

            raise ValueError(
                f"{col} pada Master masih memiliki data kosong."
            )

    return indicator_master_df


# ============================================================
# DETECT QUESTION COLUMNS
# ============================================================

def detect_question_columns(df, master_questions):

    master_question_set = {
        normalize_text(q)
        for q in master_questions
        if pd.notna(q)
    }

    question_columns = []

    for col in df.columns:

        if normalize_text(col) in master_question_set:
            question_columns.append(col)

    return question_columns


# ============================================================
# BUILD ANSWER MAPPING
# ============================================================

def build_answer_mapping(indicator_master_df):

    answer_mapping = {}

    for _, row in indicator_master_df.iterrows():

        question_key = row["question_normalized"]

        answer_mapping[question_key] = {
            normalize_text(row["level_1"]): 1,
            normalize_text(row["level_2"]): 2,
            normalize_text(row["level_3"]): 3,
            normalize_text(row["level_4"]): 4,
            normalize_text(row["level_5"]): 5,
        }

    return answer_mapping


# ============================================================
# PROCESS ASSESSMENT
# ============================================================

def process_assessment(
    df,
    question_columns,
    rater_type,
    indicator_master_df,
    answer_mapping,
):

    results = []

    for _, row in df.iterrows():

        row_dict = {
            normalize_text(k): v
            for k, v in row.items()
        }

        def field(original_name, default=""):
            return row.get(
                original_name,
                row_dict.get(normalize_text(original_name), default),
            )

        # ====================================================
        # EMPLOYEE / SUPERIOR
        # ====================================================

        if rater_type == "Self":

            employee_name = clean_value(field("Nama"))
            superior_name = ""

        else:

            employee_name = clean_value(field("Nama bawahan yang dinilai"))
            superior_name = clean_value(field("Nama"))

        # ====================================================
        # IDENTITY
        # ====================================================

        email = clean_value(field("Email"))
        division = clean_value(field("Divisi"))
        unit = clean_value(field("Unit"))
        department = clean_value(field("Departemen"))
        directorate = clean_value(field("Direktorat"))
        current_level = clean_value(field("Level Saat Ini"))
        target_level = clean_value(field("Level Tujuan"))
        next_path = clean_value(field("Level Next Path"))

        completion_time = field("Completion time", None)
        last_modified_time = field("Last modified time", None)

        base_record = {
            "rater_type": rater_type,
            "employee_name": employee_name,
            "superior_name": superior_name,
            "email": email,
            "division": division,
            "unit": unit,
            "department": department,
            "directorate": directorate,
            "current_level": current_level,
            "target_level": target_level,
            "next_path": next_path,
            "completion_time": completion_time,
            "last_modified_time": last_modified_time,
        }

        # ====================================================
        # QUESTIONS
        # ====================================================

        for question_col in question_columns:

            question_normalized = normalize_text(question_col)

            answer = row.get(question_col, "")

            if pd.isna(answer):
                continue

            answer_str = str(answer).strip()

            if not answer_str:
                continue

            answer_normalized = normalize_text(answer_str)

            # =================================================
            # FIND QUESTION
            # =================================================

            master_rows = indicator_master_df[
                indicator_master_df["question_normalized"]
                == question_normalized
            ]

            if master_rows.empty:

                record = dict(base_record)
                record.update({
                    "question": question_col,
                    "answer_text": answer_str,
                    "rating": np.nan,
                    "mapping_status": "QUESTION_NOT_FOUND",
                    "indicator_id": "",
                    "dimension": "",
                    "indicator_no": np.nan,
                    "indicator": "",
                })

                results.append(record)

                continue

            master_row = master_rows.iloc[0]

            # =================================================
            # ANSWER -> RATING
            # =================================================

            rating = answer_mapping[question_normalized].get(
                answer_normalized,
                np.nan,
            )

            mapping_status = (
                "ANSWER_NOT_FOUND" if pd.isna(rating) else "MATCHED"
            )

            record = dict(base_record)
            record.update({
                "question": question_col,
                "answer_text": answer_str,
                "rating": rating,
                "mapping_status": mapping_status,
                "indicator_id": master_row["indicator_id"],
                "dimension": master_row["dimension"],
                "indicator_no": master_row["indicator_no"],
                "indicator": master_row["indicator"],
            })

            results.append(record)

    return pd.DataFrame(results)


# ============================================================
# PREPARE EMPLOYEE VALUES
# ============================================================

def prepare_employee_values(
    employee_name,
    scored_answers,
    df_self,
    df_superior,
    indicator_master_df,
):

    target_employee = normalize_text(employee_name)

    employee_data = scored_answers[
        scored_answers["employee_name"].apply(normalize_text)
        == target_employee
    ].copy()

    if employee_data.empty:

        raise ValueError(
            "Tidak ada data assessment untuk employee: "
            f"{employee_name}"
        )

    values = {"Identitas": {}}

    # ========================================================
    # D2 — EMPLOYEE
    # ========================================================

    values["Identitas"]["D2"] = employee_name

    # ========================================================
    # SELF RAW
    # ========================================================

    self_raw = pd.DataFrame()

    if df_self is not None and not df_self.empty:

        self_name_col = find_column(df_self, ["Nama", "Name"])

        if self_name_col is not None:

            self_raw = df_self[
                df_self[self_name_col].apply(normalize_text)
                == target_employee
            ].copy()

    # ========================================================
    # SUPERIOR
    # ========================================================

    superior_name = None

    if df_superior is not None and not df_superior.empty:

        subordinate_col = find_column(
            df_superior,
            [
                "Nama bawahan yang dinilai",
                "Nama Bawahan",
                "Bawahan",
            ],
        )

        if subordinate_col is not None:

            superior_employee_data = df_superior[
                df_superior[subordinate_col].apply(normalize_text)
                == target_employee
            ].copy()

            if not superior_employee_data.empty:

                superior_name_col = find_column(
                    df_superior,
                    [
                        "Nama",
                        "Name",
                        "Nama atasan yang menilai",
                        "Nama atasan",
                        "Nama superior",
                        "Superior Name",
                    ],
                )

                if superior_name_col is not None:

                    superior_name = clean_excel_value(
                        superior_employee_data.iloc[0][superior_name_col]
                    )

    values["Identitas"]["D3"] = superior_name

    # ========================================================
    # IDENTITY MAPPING
    # ========================================================

    identity_mapping = {
        "D4": ["Level Saat Ini", "Level saat ini", "Current Level"],
        "D5": ["Level Tujuan", "Level tujuan", "Target Level"],
        "D6": [
            "Tanggal Assesment",
            "Tanggal Assessment",
            "Assessment Date",
            "Completion time",
            "Start time",
            "Last modified time",
        ],
        "D7": ["Unit"],
        "D8": ["Departemen", "Department"],
        "D9": ["Divisi", "Division"],
        "D10": ["Direktorat", "Directorate"],
        "D12": ["Level Next Path", "Next Path", "Level Next"],
    }

    # ========================================================
    # FILL IDENTITY FROM SELF
    # ========================================================

    if not self_raw.empty:

        self_row = self_raw.iloc[0]

        for target_cell, candidates in identity_mapping.items():

            source_col = find_column(self_raw, candidates)

            if source_col is None:
                continue

            value = clean_excel_value(self_row[source_col])

            if value is None:
                continue

            if target_cell == "D6":

                if isinstance(value, pd.Timestamp):
                    value = value.to_pydatetime()

            values["Identitas"][target_cell] = value

    # ========================================================
    # 5 DIMENSIONS
    # ========================================================

    for dimension_key, sheet_name in DIMENSIONS.items():

        values[sheet_name] = {}

        dimension_data = employee_data[
            employee_data["dimension"] == dimension_key
        ].copy()

        # Use the actual indicator rows from the Master for this dimension.
        # This supports different question counts per sheet, e.g.
        # People Agility = 16 and Result Agility = 12.
        dimension_master = (
            indicator_master_df[
                indicator_master_df["dimension"] == dimension_key
            ]
            .sort_values("indicator_no")
        )

        for _, master_row in dimension_master.iterrows():

            indicator_id = master_row["indicator_id"]
            row = int(master_row["master_row"])

            indicator_data = dimension_data[
                dimension_data["indicator_id"] == indicator_id
            ].copy()

            # =================================================
            # SELF
            # =================================================

            self_data = indicator_data[
                indicator_data["rater_type"] == "Self"
            ].copy()

            self_rating = None

            if not self_data.empty:

                ratings = pd.to_numeric(
                    self_data["rating"],
                    errors="coerce",
                ).dropna()

                if not ratings.empty:
                    self_rating = ratings.iloc[0]

            values[sheet_name][f"I{row}"] = clean_excel_value(self_rating)

            self_comment = None

            if not self_data.empty:

                self_answers = (
                    self_data["answer_text"]
                    .dropna()
                    .astype(str)
                    .str.strip()
                )

                self_answers = self_answers[self_answers != ""]

                if not self_answers.empty:
                    self_comment = "\n".join(self_answers.tolist())

            values[sheet_name][f"J{row}"] = clean_excel_value(self_comment)

            # =================================================
            # SUPERIOR
            # =================================================

            superior_data = indicator_data[
                indicator_data["rater_type"] == "Superior"
            ].copy()

            superior_rating = None

            if not superior_data.empty:

                ratings = pd.to_numeric(
                    superior_data["rating"],
                    errors="coerce",
                ).dropna()

                if not ratings.empty:
                    superior_rating = ratings.iloc[0]

            values[sheet_name][f"K{row}"] = clean_excel_value(
                superior_rating
            )

            superior_comment = None

            if not superior_data.empty:

                superior_answers = (
                    superior_data["answer_text"]
                    .dropna()
                    .astype(str)
                    .str.strip()
                )

                superior_answers = superior_answers[superior_answers != ""]

                if not superior_answers.empty:
                    superior_comment = "\n".join(superior_answers.tolist())

            values[sheet_name][f"L{row}"] = clean_excel_value(
                superior_comment
            )

    return values


# ============================================================
# XLSX XML WRITER  (PENGGANTI PYWIN32 / EXCEL COM)
#
# Menulis nilai langsung ke XML di dalam paket .xlsx.
# - Hanya memakai library bawaan Python (tanpa lxml / pywin32)
# - Style cell (atribut s=) dipertahankan
# - Chart, gambar, merge, validation, dll. tidak tersentuh
# - Formula dihitung ulang otomatis saat file dibuka di Excel
# ============================================================

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
NS_PKG_REL = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
NS_CONTENT_TYPES = (
    "http://schemas.openxmlformats.org/package/2006/content-types"
)
NS_XML = "http://www.w3.org/XML/1998/namespace"

_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_EXCEL_EPOCH = datetime(1899, 12, 30)
_START_TAG_RE = re.compile(rb"<[A-Za-z_][^>]*>")


def _q(tag):

    return f"{{{NS_MAIN}}}{tag}"


def _col_to_num(col_letters):

    number = 0

    for ch in col_letters.upper():
        number = number * 26 + (ord(ch) - 64)

    return number


def _split_address(address):

    match = re.fullmatch(r"([A-Za-z]+)(\d+)", address.strip())

    if not match:
        raise ValueError(f"Alamat cell tidak valid: {address}")

    return match.group(1).upper(), int(match.group(2))


def _insert_before(parent, reference, new_element):

    parent.insert(list(parent).index(reference), new_element)


def _insert_after(parent, reference, new_element):

    parent.insert(list(parent).index(reference) + 1, new_element)


# ------------------------------------------------------------
# PARSE / SERIALIZE
#
# ElementTree bawaan Python bisa mengubah prefix namespace dan
# membuang deklarasi xmlns yang tidak dipakai (mis. yang dirujuk
# oleh mc:Ignorable) -> Excel bisa minta "repair".
# Solusi: daftarkan semua namespace asli, lalu tag root asli
# dikembalikan persis seperti aslinya setelah serialisasi.
# ------------------------------------------------------------

def _parse_xml(data):

    ns_pairs = []

    for _, item in ET.iterparse(io.BytesIO(data), events=("start-ns",)):
        ns_pairs.append(item)

    for prefix, uri in ns_pairs:

        try:
            ET.register_namespace(prefix, uri)
        except ValueError:
            pass

    root = ET.fromstring(data)

    return root, ns_pairs


def _serialize_xml(root, original_bytes, ns_pairs):

    out = ET.tostring(root, encoding="UTF-8", xml_declaration=True)

    original_match = _START_TAG_RE.search(original_bytes)
    new_match = _START_TAG_RE.search(out)

    if original_match is None or new_match is None:
        return out

    root_tag = original_match.group(0)

    # pastikan semua namespace yang pernah dipakai tetap terdeklarasi
    for prefix, uri in ns_pairs:

        if prefix:
            pattern = rb"\sxmlns:" + re.escape(prefix.encode()) + rb"="
            declaration = f' xmlns:{prefix}="{uri}"'.encode()
        else:
            pattern = rb"\sxmlns="
            declaration = f' xmlns="{uri}"'.encode()

        if re.search(pattern, root_tag) is None:

            if root_tag.endswith(b"/>"):
                root_tag = root_tag[:-2] + declaration + b"/>"
            else:
                root_tag = root_tag[:-1] + declaration + b">"

    return out[:new_match.start()] + root_tag + out[new_match.end():]


def _normalize_value(value):

    if value is None:
        return None

    if isinstance(value, np.generic):
        value = value.item()

    if isinstance(value, float) and math.isnan(value):
        return None

    if isinstance(value, pd.Timestamp):

        if pd.isna(value):
            return None

        value = value.to_pydatetime()

    return value


def _to_excel_serial(value):

    if isinstance(value, datetime):

        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)

    else:

        value = datetime(value.year, value.month, value.day)

    delta = value - _EXCEL_EPOCH

    return delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6


def _set_cell_value(cell, value):
    """
    Isi satu elemen <c>. Atribut r (alamat) dan s (style) dipertahankan.
    Formula lama pada cell dihapus (sama seperti Range.Value di COM).
    """

    for child in list(cell):
        cell.remove(child)

    for attr in list(cell.attrib):
        if attr not in ("r", "s"):
            del cell.attrib[attr]

    value = _normalize_value(value)

    # ---------------------------------------------- kosong
    if value is None:
        return

    # ---------------------------------------------- boolean
    if isinstance(value, bool):

        cell.set("t", "b")

        v = ET.SubElement(cell, _q("v"))
        v.text = "1" if value else "0"

        return

    # ---------------------------------------------- angka
    if isinstance(value, (int, float)):

        v = ET.SubElement(cell, _q("v"))

        if isinstance(value, float) and value.is_integer():
            v.text = str(int(value))
        else:
            v.text = repr(value)

        return

    # ---------------------------------------------- tanggal
    if isinstance(value, (datetime, date)):

        v = ET.SubElement(cell, _q("v"))
        v.text = repr(_to_excel_serial(value))

        return

    # ---------------------------------------------- teks
    text = _ILLEGAL_XML_CHARS.sub("", str(value))
    text = text[:32767]

    cell.set("t", "inlineStr")

    is_el = ET.SubElement(cell, _q("is"))
    t_el = ET.SubElement(is_el, _q("t"))
    t_el.set(f"{{{NS_XML}}}space", "preserve")
    t_el.text = text


def _get_or_create_row(sheet_data, row_num):

    for row in sheet_data.findall(_q("row")):

        r_attr = row.get("r")

        if r_attr is None:
            continue

        r_num = int(r_attr)

        if r_num == row_num:
            return row

        if r_num > row_num:

            new_row = ET.Element(_q("row"))
            new_row.set("r", str(row_num))
            _insert_before(sheet_data, row, new_row)

            return new_row

    new_row = ET.SubElement(sheet_data, _q("row"))
    new_row.set("r", str(row_num))

    return new_row


def _get_or_create_cell(row, address):

    col_letters, _ = _split_address(address)
    target_col = _col_to_num(col_letters)

    for cell in row.findall(_q("c")):

        c_addr = cell.get("r")

        if c_addr is None:
            continue

        c_col = _col_to_num(_split_address(c_addr)[0])

        if c_col == target_col:
            return cell

        if c_col > target_col:

            new_cell = ET.Element(_q("c"))
            new_cell.set("r", address)
            _insert_before(row, cell, new_cell)

            return new_cell

    new_cell = ET.SubElement(row, _q("c"))
    new_cell.set("r", address)

    return new_cell


def _write_cells_to_sheet_xml(sheet_xml_bytes, cell_values):

    root, ns_pairs = _parse_xml(sheet_xml_bytes)

    sheet_data = root.find(_q("sheetData"))

    if sheet_data is None:
        sheet_data = ET.SubElement(root, _q("sheetData"))

    for address, value in cell_values.items():

        _, row_num = _split_address(address)

        row = _get_or_create_row(sheet_data, row_num)

        # atribut spans (opsional) bisa jadi tidak akurat lagi
        if "spans" in row.attrib:
            del row.attrib["spans"]

        cell = _get_or_create_cell(row, address)

        _set_cell_value(cell, value)

    return _serialize_xml(root, sheet_xml_bytes, ns_pairs)


def _resolve_sheet_paths(contents):
    """
    Peta: nama sheet -> path XML di dalam paket .xlsx
    """

    wb_root = ET.fromstring(contents["xl/workbook.xml"])
    rels_root = ET.fromstring(contents["xl/_rels/workbook.xml.rels"])

    rel_map = {
        rel.get("Id"): rel.get("Target")
        for rel in rels_root.findall(f"{{{NS_PKG_REL}}}Relationship")
    }

    sheet_paths = {}

    sheets_el = wb_root.find(_q("sheets"))

    for sheet in sheets_el.findall(_q("sheet")):

        name = sheet.get("name")
        rid = sheet.get(f"{{{NS_REL}}}id")
        target = rel_map.get(rid)

        if target is None:
            continue

        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = posixpath.normpath(posixpath.join("xl", target))

        sheet_paths[name] = path

    return sheet_paths


def _force_full_recalc(contents):
    """
    - set fullCalcOnLoad => Excel menghitung ulang semua formula
      (chart & skor ikut ter-update) saat file dibuka
    - hapus calcChain (akan dibangun ulang oleh Excel) agar tidak
      terjadi pesan 'repair' setelah cell formula diganti nilai
    """

    # ---------------- workbook.xml
    original_wb = contents["xl/workbook.xml"]

    wb_root, wb_ns = _parse_xml(original_wb)

    calc_pr = wb_root.find(_q("calcPr"))

    if calc_pr is None:

        calc_pr = ET.Element(_q("calcPr"))

        anchor = None

        for tag in [
            "sheets",
            "functionGroups",
            "externalReferences",
            "definedNames",
        ]:
            found = wb_root.find(_q(tag))

            if found is not None:
                anchor = found

        if anchor is not None:
            _insert_after(wb_root, anchor, calc_pr)
        else:
            wb_root.append(calc_pr)

    calc_pr.set("fullCalcOnLoad", "1")

    contents["xl/workbook.xml"] = _serialize_xml(
        wb_root,
        original_wb,
        wb_ns,
    )

    # ---------------- calcChain
    if "xl/calcChain.xml" in contents:

        del contents["xl/calcChain.xml"]

        original_rels = contents["xl/_rels/workbook.xml.rels"]

        rels_root, rels_ns = _parse_xml(original_rels)

        for rel in rels_root.findall(f"{{{NS_PKG_REL}}}Relationship"):

            if (rel.get("Type") or "").endswith("/calcChain"):
                rels_root.remove(rel)

        contents["xl/_rels/workbook.xml.rels"] = _serialize_xml(
            rels_root,
            original_rels,
            rels_ns,
        )

        if "[Content_Types].xml" in contents:

            original_ct = contents["[Content_Types].xml"]

            ct_root, ct_ns = _parse_xml(original_ct)

            for override in ct_root.findall(
                f"{{{NS_CONTENT_TYPES}}}Override"
            ):

                if override.get("PartName") == "/xl/calcChain.xml":
                    ct_root.remove(override)

            contents["[Content_Types].xml"] = _serialize_xml(
                ct_root,
                original_ct,
                ct_ns,
            )


def write_values_to_xlsx(output_file, employee_values):
    """
    Pengganti write_values_with_excel() (pywin32 / Excel COM).
    Menulis employee_values ke file hasil copy Master.
    """

    output_file = Path(output_file)

    try:

        with zipfile.ZipFile(output_file, "r") as zin:

            infos = zin.infolist()
            contents = {info.filename: zin.read(info.filename) for info in infos}

        order = [info.filename for info in infos]

        sheet_paths = _resolve_sheet_paths(contents)

        lower_sheet_paths = {k.lower(): v for k, v in sheet_paths.items()}

        # ----------------------------------------------------
        # WRITE VALUES PER SHEET
        # ----------------------------------------------------

        for sheet_name, cell_values in employee_values.items():

            path = sheet_paths.get(sheet_name) or lower_sheet_paths.get(
                sheet_name.lower()
            )

            if path is None or path not in contents:

                raise ValueError(
                    f"Sheet '{sheet_name}' tidak ditemukan di file Master."
                )

            contents[path] = _write_cells_to_sheet_xml(
                contents[path],
                cell_values,
            )

        # ----------------------------------------------------
        # FORCE RECALC (setara workbook.Save() pada Excel)
        # ----------------------------------------------------

        _force_full_recalc(contents)

        # ----------------------------------------------------
        # SAVE (tulis ke temp lalu replace, agar atomic)
        # ----------------------------------------------------

        tmp_file = output_file.with_suffix(".tmp")

        with zipfile.ZipFile(
            tmp_file,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as zout:

            for name in order:

                if name in contents:
                    zout.writestr(name, contents[name])

        os.replace(tmp_file, output_file)

    except Exception as e:

        raise RuntimeError(
            "Gagal menulis data ke Excel.\n\n"
            f"Detail: {e}"
        ) from e


# ============================================================
# VALIDATE OUTPUT
# ============================================================

def validate_employee_output(output_file, employee_name):

    wb = load_workbook(output_file, data_only=False)

    errors = []

    try:

        if not wb.sheetnames:

            errors.append("Workbook tidak memiliki sheet.")

        elif wb.sheetnames[0] != "Identitas":

            errors.append("Sheet Identitas tidak sesuai.")

        if "Identitas" not in wb.sheetnames:

            errors.append("Sheet Identitas tidak ditemukan.")

        else:

            output_name = wb["Identitas"]["D2"].value

            if (
                normalize_text(output_name)
                != normalize_text(employee_name)
            ):

                errors.append("Identitas!D2 tidak sesuai employee.")

        for sheet_name in ASSESSMENT_SHEETS:

            if sheet_name not in wb.sheetnames:

                errors.append(f"Sheet {sheet_name} tidak ditemukan.")

    finally:

        wb.close()

    if errors:
        return {"status": "FAILED", "errors": errors}

    return {"status": "PASSED", "errors": []}


# ============================================================
# CREATE ONE EMPLOYEE FILE
# ============================================================

def create_employee_file(
    employee_name,
    master_file,
    df_self,
    df_superior,
    scored_answers,
    indicator_master_df,
    output_dir,
):

    output_dir.mkdir(parents=True, exist_ok=True)

    safe_name = re.sub(r'[\\/*?:"<>|]', "_", str(employee_name)).strip()

    if not safe_name:
        safe_name = "Employee"

    output_file = output_dir / f"KompetensixLearning agility assesment_{safe_name}.xlsx"

    # ========================================================
    # STEP 1 — COPY MASTER
    # ========================================================

    shutil.copy2(master_file, output_file)

    # ========================================================
    # STEP 2 — PREPARE DATA
    # ========================================================

    employee_values = prepare_employee_values(
        employee_name=employee_name,
        scored_answers=scored_answers,
        df_self=df_self,
        df_superior=df_superior,
        indicator_master_df=indicator_master_df,
    )

    # ========================================================
    # STEP 3 — WRITE VALUES (XLSX XML, tanpa Excel COM)
    # ========================================================

    write_values_to_xlsx(
        output_file=output_file,
        employee_values=employee_values,
    )

    # ========================================================
    # STEP 4 — BASIC CHECK
    # ========================================================

    if not output_file.exists():

        raise FileNotFoundError(
            f"Output tidak berhasil dibuat:\n{output_file}"
        )

    if output_file.stat().st_size == 0:

        raise IOError(f"Output file kosong:\n{output_file}")

    # ========================================================
    # STEP 5 — VALIDATION
    # ========================================================

    validation = validate_employee_output(
        output_file=output_file,
        employee_name=employee_name,
    )

    if validation["status"] != "PASSED":
        raise ValueError(str(validation["errors"]))

    return output_file


# ============================================================
# CREATE ALL EMPLOYEE FILES
# ============================================================

def create_all_employee_files(
    employees,
    master_file,
    df_self,
    df_superior,
    scored_answers,
    indicator_master_df,
    output_dir,
):

    output_dir.mkdir(parents=True, exist_ok=True)

    generated_files = []

    for employee_name in employees:

        output_file = create_employee_file(
            employee_name=employee_name,
            master_file=master_file,
            df_self=df_self,
            df_superior=df_superior,
            scored_answers=scored_answers,
            indicator_master_df=indicator_master_df,
            output_dir=output_dir,
        )

        generated_files.append(output_file)

    # ========================================================
    # CREATE ZIP
    # ========================================================

    zip_file = output_dir / "Learning_Agility_All_Employees.zip"

    if zip_file.exists():
        zip_file.unlink()

    with zipfile.ZipFile(
        zip_file,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zipf:

        for output_file in generated_files:

            if not output_file.exists():

                raise FileNotFoundError(
                    "File employee tidak ditemukan "
                    "sebelum dimasukkan ke ZIP:\n"
                    f"{output_file}"
                )

            zipf.write(output_file, arcname=output_file.name)

    # ========================================================
    # VALIDATE ZIP
    # ========================================================

    if not zip_file.exists():
        raise FileNotFoundError("ZIP tidak berhasil dibuat.")

    if zip_file.stat().st_size == 0:
        raise IOError("ZIP berhasil dibuat tetapi kosong.")

    return zip_file, generated_files


# ============================================================
# PROCESS ALL DATA
# ============================================================

def process_files(master_file, self_file, superior_file):

    if not master_file.exists():

        raise FileNotFoundError(
            "Master Template tidak ditemukan.\n\n"
            f"Path:\n{master_file}"
        )

    if hasattr(self_file, "seek"):
        self_file.seek(0)

    if hasattr(superior_file, "seek"):
        superior_file.seek(0)

    df_self = pd.read_excel(self_file, engine="openpyxl")

    df_superior = pd.read_excel(superior_file, engine="openpyxl")

    indicator_master_df = read_master_indicators(master_file)

    master_questions = indicator_master_df["question"].tolist()

    self_question_cols = detect_question_columns(df_self, master_questions)

    superior_question_cols = detect_question_columns(
        df_superior,
        master_questions,
    )

    expected_question_count = len(indicator_master_df)

    if len(self_question_cols) != expected_question_count:

        raise ValueError(
            "Self raw contains "
            f"{len(self_question_cols)} questions. "
            f"Expected {expected_question_count} based on Master."
        )

    if len(superior_question_cols) != expected_question_count:

        raise ValueError(
            "Superior raw contains "
            f"{len(superior_question_cols)} questions. "
            f"Expected {expected_question_count} based on Master."
        )

    # Ensure every Master question exists in both RAW exports.
    # detect_question_columns() already matches by normalized question text,
    # so this also catches missing/new questions without relying on column order.
    master_question_set = set(indicator_master_df["question_normalized"])
    self_question_set = {normalize_text(q) for q in self_question_cols}
    superior_question_set = {normalize_text(q) for q in superior_question_cols}

    missing_self = sorted(master_question_set - self_question_set)
    missing_superior = sorted(master_question_set - superior_question_set)

    if missing_self:
        raise ValueError(
            "Ada question Master yang tidak ditemukan di Self RAW: "
            + "; ".join(missing_self[:5])
            + (" ..." if len(missing_self) > 5 else "")
        )

    if missing_superior:
        raise ValueError(
            "Ada question Master yang tidak ditemukan di Superior RAW: "
            + "; ".join(missing_superior[:5])
            + (" ..." if len(missing_superior) > 5 else "")
        )

    answer_mapping = build_answer_mapping(indicator_master_df)

    self_scored_df = process_assessment(
        df=df_self,
        question_columns=self_question_cols,
        rater_type="Self",
        indicator_master_df=indicator_master_df,
        answer_mapping=answer_mapping,
    )

    superior_scored_df = process_assessment(
        df=df_superior,
        question_columns=superior_question_cols,
        rater_type="Superior",
        indicator_master_df=indicator_master_df,
        answer_mapping=answer_mapping,
    )

    scored_answers_df = pd.concat(
        [self_scored_df, superior_scored_df],
        ignore_index=True,
    )

    empty_employee = (
        scored_answers_df["employee_name"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("")
        .sum()
    )

    if empty_employee > 0:

        raise ValueError(
            f"Ada {empty_employee} record "
            "dengan employee_name kosong."
        )

    employees = sorted(
        scored_answers_df[
            scored_answers_df["rater_type"] == "Self"
        ]["employee_name"]
        .dropna()
        .astype(str)
        .str.strip()
        .loc[lambda s: s != ""]
        .unique()
    )

    if len(employees) == 0:

        raise ValueError(
            "Tidak ada employee dari Self Assessment."
        )

    return {
        "df_self": df_self,
        "df_superior": df_superior,
        "indicator_master_df": indicator_master_df,
        "self_scored_df": self_scored_df,
        "superior_scored_df": superior_scored_df,
        "scored_answers_df": scored_answers_df,
        "employees": employees,
    }


# ============================================================
# SESSION STATE
# ============================================================

_STATE_DEFAULTS = {
    "processed_data": None,
    "generated_file": None,
    "generated_employee": None,
    "generated_zip": None,
    "self_file_version": 0,
    "superior_file_version": 0,
    "self_file_bytes": None,
    "self_file_name": None,
    "superior_file_bytes": None,
    "superior_file_name": None,
    "mapping_counts": {},
}

for _key, _default in _STATE_DEFAULTS.items():

    if _key not in st.session_state:
        st.session_state[_key] = _default


# ============================================================
# BACKGROUND
# ============================================================

if BACKGROUND_BASE64:

    st.markdown(
        f"""
        <style>
        .stApp {{
            min-height: 100vh;
            background-image:
                linear-gradient(
                    rgba(255,248,242,0.68),
                    rgba(255,244,235,0.76)
                ),
                url("data:image/png;base64,{BACKGROUND_BASE64}");
            background-size: cover;
            background-position: center center;
            background-repeat: no-repeat;
            background-attachment: fixed;
        }}
        </style>
        """,
        unsafe_allow_html=True,
    )

else:

    st.markdown(
        """
        <style>
        .stApp {
            min-height: 100vh;
            background: linear-gradient(
                135deg,
                #fffaf7 0%,
                #fff4eb 48%,
                #fffaf7 100%
            );
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    textwrap.dedent(
        """
        <style>

        /* STREAMLIT STATUS / SPINNER TEXT */
        div[data-testid="stSpinner"] { color: #000000 !important; }
        div[data-testid="stSpinner"] p { color: #000000 !important; }

        /* SELECTBOX */
        div[data-testid="stSelectbox"] label { color: #000000 !important; }
        div[data-baseweb="select"] { color: #000000 !important; }
        div[data-baseweb="select"] * { color: #000000 !important; }

        /* MAIN AREA */
        .block-container {
            max-width: 820px !important;
            padding-top: 30px !important;
            padding-bottom: 55px !important;
            position: relative;
            z-index: 2;
        }

        /* HEADER */
        header[data-testid="stHeader"] { background: transparent; }
        #MainMenu { visibility: hidden; }
        footer { visibility: hidden; }

        /* REMOVE OLD WAVES */
        .stApp::before, .stApp::after { display: none; }

        /* MAIN CARD */
        .st-key-main_card {
            width: min(700px,100%);
            margin: 0 auto;
            padding: 26px 42px 25px 42px;
            background: rgba(255,255,255,0.90);
            border: 1px solid rgba(255,155,83,0.37);
            border-radius: 26px;
            box-shadow:
                0 25px 60px rgba(211,125,68,0.15),
                0 7px 22px rgba(211,125,68,0.08);
            backdrop-filter: blur(18px);
            -webkit-backdrop-filter: blur(18px);
        }

        /* BRAND */
        .learning-brand { text-align: center; margin-bottom: 18px; }

        .learning-logo {
            width: 54px; height: 54px;
            margin: 0 auto 5px auto;
            position: relative;
        }

        .logo-piece-one {
            position: absolute; left: 8px; top: 3px;
            width: 24px; height: 46px;
            background: #ff7417;
            clip-path: polygon(38% 0%, 100% 0%, 62% 34%, 62% 100%, 0% 78%, 0% 32%);
        }

        .logo-piece-two {
            position: absolute; left: 24px; top: 7px;
            width: 25px; height: 38px;
            background: #f87618;
            clip-path: polygon(26% 0%, 100% 0%, 100% 56%, 52% 82%, 0% 100%, 27% 60%);
        }

        .logo-piece-three {
            position: absolute; left: 8px; bottom: 3px;
            width: 24px; height: 18px;
            background: #ff9635;
            clip-path: polygon(0% 0%, 100% 45%, 50% 100%);
        }

        .learning-brand-title {
            margin: 0;
            color: #192536;
            font-size: 34px;
            line-height: 1.08;
            font-weight: 750;
            letter-spacing: -1.3px;
        }

        .learning-brand-title .orange { color: #f87618; }

        .learning-brand-subtitle {
            margin-top: 6px;
            color: #7c8797;
            font-size: 15px;
            line-height: 1.3;
        }

        /* SECTION LABEL */
        .learning-section-label {
            display: flex;
            align-items: center;
            gap: 10px;
            margin-top: 15px;
            margin-bottom: 8px;
            color: #172333;
            font-size: 17px;
            font-weight: 650;
        }

        /* ICON */
        .user-icon, .people-icon {
            width: 29px; height: 29px; min-width: 29px;
            border-radius: 8px;
            background: rgba(255,116,24,0.09);
            position: relative;
        }

        .user-icon::before {
            content: ""; position: absolute;
            width: 7px; height: 7px; left: 9px; top: 5px;
            border: 2px solid #f87618; border-radius: 50%;
        }

        .user-icon::after {
            content: ""; position: absolute;
            width: 14px; height: 8px; left: 6px; bottom: 5px;
            border: 2px solid #f87618; border-bottom: 0;
            border-radius: 8px 8px 0 0;
        }

        .people-icon::before {
            content: ""; position: absolute;
            width: 6px; height: 6px; left: 5px; top: 5px;
            border: 2px solid #f87618; border-radius: 50%;
            box-shadow: 10px 2px 0 -1px #ffffff, 10px 2px 0 1px #f87618;
        }

        .people-icon::after {
            content: ""; position: absolute;
            width: 10px; height: 7px; left: 4px; bottom: 5px;
            border: 2px solid #f87618; border-bottom: 0;
            border-radius: 8px 8px 0 0;
            box-shadow: 10px 0 0 -1px #ffffff, 10px 0 0 1px #f87618;
        }

        /* UPLOAD ZONE WRAPPER */
        .st-key-self_upload_zone, .st-key-superior_upload_zone {
            position: relative; width: 100%; margin: 0; padding: 0;
        }

        /* CUSTOM EMPTY UPLOAD CARD */
        .custom-empty-upload-card {
            width: 100%; height: 96px; min-height: 96px;
            box-sizing: border-box;
            display: flex; align-items: center;
            position: relative;
            padding: 20px 26px;
            background: rgba(255,255,255,0.86);
            border: 1px solid rgba(255,137,76,0.36);
            border-radius: 17px;
            box-shadow: 0 5px 18px rgba(213,125,67,0.035);
            transition: transform 0.18s ease, box-shadow 0.18s ease, border-color 0.18s ease;
            pointer-events: none;
        }

        .st-key-self_upload_zone:hover .custom-empty-upload-card,
        .st-key-superior_upload_zone:hover .custom-empty-upload-card {
            transform: translateY(-1px);
            box-shadow: 0 14px 28px rgba(255,116,18,0.18);
            border-color: rgba(255,126,49,0.50);
        }

        /* UPLOAD ICON BOX */
        .upload-icon-box {
            width: 62px; height: 62px; min-width: 62px;
            margin-right: 20px;
            display: flex; align-items: center; justify-content: center;
            position: relative;
            border-radius: 13px;
            background: #fff0e6;
        }

        .upload-arrow { width: 26px; height: 31px; position: relative; }

        .upload-arrow::before {
            content: ""; position: absolute;
            width: 2px; height: 23px; left: 12px; top: 6px;
            background: #ff7417; border-radius: 2px;
        }

        .upload-arrow::after {
            content: ""; position: absolute;
            width: 11px; height: 11px; left: 7px; top: 4px;
            border-top: 2px solid #ff7417;
            border-left: 2px solid #ff7417;
            transform: rotate(45deg);
        }

        /* EMPTY UPLOAD TEXT */
        .upload-text-area { flex: 1; min-width: 0; }

        .upload-primary-text {
            color: #ff7417; font-size: 20px; line-height: 1.2; font-weight: 700;
        }

        .upload-secondary-text {
            margin-top: 7px; color: #8290a2; font-size: 15px; line-height: 1.3;
        }

        /* FILE OUTLINE ICON */
        .upload-file-icon {
            width: 31px; height: 25px; min-width: 31px;
            position: relative; margin-left: 16px;
            border: 2px solid #ff7417;
            transform: skewX(-10deg);
        }

        .upload-file-icon::before {
            content: ""; position: absolute;
            width: 11px; height: 11px; right: -2px; top: -2px;
            background: rgba(255,255,255,0.90);
            border-left: 2px solid #ff7417;
            border-bottom: 2px solid #ff7417;
        }

        /* NATIVE STREAMLIT UPLOADER (hidden visually, still clickable) */
        .st-key-self_native_uploader, .st-key-superior_native_uploader {
            position: absolute !important;
            left: 0 !important; top: 0 !important;
            width: 100% !important;
            height: 126px !important; min-height: 126px !important;
            z-index: 50 !important;
            opacity: 0 !important;
            margin: 0 !important; padding: 0 !important;
        }

        .st-key-self_native_uploader div[data-testid="stFileUploader"],
        .st-key-superior_native_uploader div[data-testid="stFileUploader"] {
            width: 100% !important;
            height: 126px !important; min-height: 126px !important;
            margin: 0 !important; padding: 0 !important;
        }

        .st-key-self_native_uploader div[data-testid="stFileUploader"] section,
        .st-key-superior_native_uploader div[data-testid="stFileUploader"] section {
            width: 100% !important;
            height: 126px !important; min-height: 126px !important;
            box-sizing: border-box !important;
            padding: 0 !important; margin: 0 !important;
            border: none !important;
            background: transparent !important;
            cursor: pointer !important;
        }

        .st-key-self_native_uploader div[data-testid="stFileUploader"] *,
        .st-key-superior_native_uploader div[data-testid="stFileUploader"] * {
            cursor: pointer !important;
        }

        /* UPLOADED FILE CARD */
        .custom-uploaded-file {
            width: 100%; height: 108px; min-height: 108px;
            box-sizing: border-box;
            display: flex; align-items: center;
            padding: 14px 18px;
            background: rgba(255,255,255,0.88);
            border: 1px solid rgba(255,151,82,0.35);
            border-radius: 15px;
            box-shadow: 0 7px 20px rgba(213,125,67,0.07);
        }

        .custom-file-badge {
            width: 58px; height: 58px; min-width: 58px;
            margin-right: 16px;
            display: flex; align-items: center; justify-content: center;
            border-radius: 11px;
            background: #ffe8d8;
            color: #f87618;
            font-size: 16px; font-weight: 750;
        }

        .custom-file-info { flex: 1; min-width: 0; }

        .custom-file-name {
            color: #172333; font-size: 15px; font-weight: 650; line-height: 1.3;
            white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
        }

        .custom-file-description {
            margin-top: 5px; color: #8a95a4; font-size: 13px;
        }

        .custom-file-check {
            width: 40px; height: 40px; min-width: 40px;
            margin-left: 16px;
            display: flex; align-items: center; justify-content: center;
            border-radius: 50%;
            background: #16a765;
            color: #ffffff;
            font-size: 22px; font-weight: 700;
        }

        /* CHANGE FILE */
        .st-key-change_self_file, .st-key-change_superior_file {
            display: flex; justify-content: flex-end;
            margin-top: 1px; margin-bottom: 0;
        }

        .st-key-change_self_file button, .st-key-change_superior_file button {
            min-height: 26px; padding: 1px 8px;
            border: none; background: transparent;
            color: #f87618; font-size: 12px; font-weight: 600;
        }

        .st-key-change_self_file button:hover,
        .st-key-change_superior_file button:hover {
            color: #d95f0f;
            background: rgba(255,116,24,0.05);
        }

        /* GENERATE BUTTON */
        .st-key-generate_report { margin-top: 15px; }

        .st-key-generate_report div.stButton > button {
            width: 100%; min-height: 54px;
            border: none; border-radius: 13px;
            background: linear-gradient(90deg, #ff7112 0%, #ff8d22 55%, #ff9d39 100%);
            color: #ffffff;
            font-size: 16px; font-weight: 750;
            box-shadow: 0 10px 25px rgba(255,116,18,0.23);
        }

        .st-key-generate_report div.stButton > button:hover {
            color: #ffffff;
            transform: translateY(-1px);
            box-shadow: 0 14px 28px rgba(255,116,18,0.29);
        }

        /* SECURITY NOTE */
        .learning-security-note {
            text-align: center; margin-top: 12px;
            color: #929aa7; font-size: 12px;
        }

        /* RESULT CARD */
        .learning-result-card {
            width: min(700px,100%);
            margin: 16px auto 0 auto;
            padding: 20px 23px;
            background: rgba(255,255,255,0.84);
            border: 1px solid rgba(255,151,75,0.22);
            border-radius: 18px;
            box-shadow: 0 10px 30px rgba(213,125,67,0.08);
            backdrop-filter: blur(12px);
            -webkit-backdrop-filter: blur(12px);
        }

        .learning-result-title {
            margin: 0; color: #182536; font-size: 20px; font-weight: 700;
        }

        .learning-result-description {
            margin-top: 6px; color: #7b8695; font-size: 14px; line-height: 1.45;
        }

        /* SELECTBOX */
        div[data-baseweb="select"] > div {
            min-height: 46px;
            border-radius: 11px;
            border-color: rgba(255,126,49,0.24);
            background: rgba(255,255,255,0.84);
        }

        /* GENERATE EMPLOYEE BUTTONS */
        div[data-testid="stHorizontalBlock"] { gap: 14px; }

        /* DOWNLOAD BUTTON */
        div[data-testid="stDownloadButton"] button {
            width: 100%; min-height: 50px;
            border-radius: 12px;
            border: 1px solid rgba(255,126,49,0.28);
            background: #ffffff;
            color: #f87618;
            font-weight: 650;
        }

        div[data-testid="stDownloadButton"] button:hover {
            border-color: rgba(255,126,49,0.48);
            color: #d95f0f;
            background: #fffaf7;
        }

        /* ALERT */
        div[data-testid="stAlert"] { border-radius: 13px; }

        /* MOBILE */
        @media (max-width: 768px) {

            .block-container {
                max-width: 100% !important;
                padding: 18px 14px 35px 14px !important;
            }

            .st-key-main_card {
                width: 100%;
                padding: 22px 19px 22px 19px;
                border-radius: 22px;
            }

            .learning-brand { margin-bottom: 15px; }

            .learning-brand-title { font-size: 29px; letter-spacing: -1px; }

            .learning-brand-subtitle { font-size: 14px; }

            .learning-section-label { font-size: 16px; margin-top: 13px; }

            .custom-empty-upload-card {
                height: 116px; min-height: 116px; padding: 17px 18px;
            }

            .st-key-self_native_uploader,
            .st-key-superior_native_uploader {
                height: 116px !important; min-height: 116px !important;
            }

            .st-key-self_native_uploader div[data-testid="stFileUploader"],
            .st-key-superior_native_uploader div[data-testid="stFileUploader"] {
                height: 116px !important; min-height: 116px !important;
            }

            .st-key-self_native_uploader div[data-testid="stFileUploader"] section,
            .st-key-superior_native_uploader div[data-testid="stFileUploader"] section {
                height: 116px !important; min-height: 116px !important;
            }

            .upload-icon-box {
                width: 52px; height: 52px; min-width: 52px; margin-right: 13px;
            }

            .upload-primary-text { font-size: 17px; }

            .upload-secondary-text { font-size: 13px; }

            .upload-file-icon {
                width: 27px; height: 22px; min-width: 27px; margin-left: 10px;
            }

            .custom-uploaded-file {
                height: 96px; min-height: 96px; padding: 12px;
            }

            .custom-file-badge {
                width: 48px; height: 48px; min-width: 48px;
                margin-right: 11px; font-size: 14px;
            }

            .custom-file-check {
                width: 34px; height: 34px; min-width: 34px;
                margin-left: 10px; font-size: 18px;
            }

            .custom-file-name { font-size: 14px; }

            .custom-file-description { font-size: 12px; }

            .learning-result-card { width: 100%; }
        }

        </style>
        """
    ),
    unsafe_allow_html=True,
)


# ============================================================
# UPLOAD SECTION (dipakai untuk Self & Superior)
# ============================================================

def render_upload_section(role, section_icon_class, section_title):

    render_html(
        f"""
        <div class="learning-section-label">
            <div class="{section_icon_class}"></div>
            <span>{section_title}</span>
        </div>
        """
    )

    stored_file = get_stored_file(role)

    version = st.session_state[f"{role}_file_version"]

    if stored_file is None:

        widget_key = f"{role}_file_{version}"

        with st.container(key=f"{role}_upload_zone"):

            render_html(
                """
                <div class="custom-empty-upload-card">
                    <div class="upload-icon-box">
                        <div class="upload-arrow"></div>
                    </div>
                    <div class="upload-text-area">
                        <div class="upload-primary-text">Choose File</div>
                        <div class="upload-secondary-text">Excel file (.xlsx)</div>
                    </div>
                    <div class="upload-file-icon"></div>
                </div>
                """
            )

            with st.container(key=f"{role}_native_uploader"):

                uploaded = st.file_uploader(
                    section_title,
                    type=["xlsx"],
                    key=widget_key,
                    label_visibility="collapsed",
                )

        if uploaded is not None:

            store_uploaded_file(role, uploaded)

            st.rerun()

    else:

        render_html(
            f"""
            <div class="custom-uploaded-file">
                <div class="custom-file-badge">XLS</div>
                <div class="custom-file-info">
                    <div class="custom-file-name">
                        {html_escape(stored_file.name)}
                    </div>
                    <div class="custom-file-description">
                        {section_title}
                    </div>
                </div>
                <div class="custom-file-check">✓</div>
            </div>
            """
        )

        with st.container(key=f"change_{role}_file"):

            change_clicked = st.button(
                "Change file",
                key=f"change_{role}_file_{version}",
            )

        if change_clicked:

            clear_uploaded_file(role)

            st.session_state.processed_data = None
            st.session_state.generated_file = None
            st.session_state.generated_employee = None
            st.session_state.generated_zip = None
            st.session_state.mapping_counts = {}

            st.session_state[f"{role}_file_version"] += 1

            st.rerun()

    return get_stored_file(role)


# ============================================================
# MAIN CARD
# ============================================================

with st.container(key="main_card"):

    render_html(
        """
        <div class="learning-brand">
            <div class="learning-logo">
                <span class="logo-piece-one"></span>
                <span class="logo-piece-two"></span>
                <span class="logo-piece-three"></span>
            </div>
            <div class="learning-brand-title">
                Learning <span class="orange">Agility</span>
            </div>
            <div class="learning-brand-subtitle">
                Learning Agility Assessment System
            </div>
        </div>
        """
    )

    self_file = render_upload_section(
        "self",
        "user-icon",
        "Self Assessment",
    )

    superior_file = render_upload_section(
        "superior",
        "people-icon",
        "Superior Assessment",
    )

    # ========================================================
    # GENERATE REPORT
    # ========================================================

    with st.container(key="generate_report"):

        generate_report = st.button(
            "✦  Generate Report",
            type="primary",
            use_container_width=True,
        )

    render_html(
        """
        <div class="learning-security-note">
            ♢&nbsp;&nbsp;
            Pastikan kedua file berformat Excel (.xlsx)
        </div>
        """
    )


# ============================================================
# PROCESS / VALIDATE
# ============================================================

if generate_report:

    if not MASTER_FILE.exists():

        st.error(
            "Master Template tidak ditemukan "
            "di folder data/master/."
        )

    elif self_file is None or superior_file is None:

        st.error(
            "Upload Self Assessment dan "
            "Superior Assessment terlebih dahulu."
        )

    else:

        with st.spinner("Generating Learning Agility Report..."):

            try:

                self_file.seek(0)
                superior_file.seek(0)

                result = process_files(
                    master_file=MASTER_FILE,
                    self_file=self_file,
                    superior_file=superior_file,
                )

                st.session_state.processed_data = result
                st.session_state.generated_file = None
                st.session_state.generated_employee = None
                st.session_state.generated_zip = None

                mapping_counts = (
                    result["scored_answers_df"]["mapping_status"]
                    .value_counts()
                    .to_dict()
                )

                st.session_state.mapping_counts = mapping_counts

                st.success("Assessment berhasil diproses.")

            except Exception as e:

                st.session_state.processed_data = None
                st.session_state.generated_file = None
                st.session_state.generated_employee = None
                st.session_state.generated_zip = None
                st.session_state.mapping_counts = {}

                st.error(f"Proses gagal: {e}")


# ============================================================
# RESULT / GENERATION AREA
# ============================================================

if st.session_state.processed_data is not None:

    data = st.session_state.processed_data

    employees = data["employees"]

    scored_answers_df = data["scored_answers_df"]

    # ========================================================
    # MAPPING WARNINGS
    # ========================================================

    mapping_counts = st.session_state.get("mapping_counts", {})

    if mapping_counts.get("ANSWER_NOT_FOUND", 0) > 0:

        st.warning(
            "Ada jawaban yang tidak cocok "
            "dengan level 1–5 pada Master."
        )

    if mapping_counts.get("QUESTION_NOT_FOUND", 0) > 0:

        st.warning(
            "Ada question yang tidak ditemukan "
            "pada Master."
        )

    # ========================================================
    # GENERATE EMPLOYEE CARD
    # ========================================================

    render_html(
        """
        <div class="learning-result-card">
            <div class="learning-result-title">
                Generate Employee Report
            </div>
            <div class="learning-result-description">
                Pilih employee untuk membuat report
                individual atau generate seluruh employee.
            </div>
        </div>
        """
    )

    selected_employee = st.selectbox("Pilih Employee", employees)

    col_generate_one, col_generate_all = st.columns(2)

    # ========================================================
    # GENERATE ONE EMPLOYEE
    # ========================================================

    with col_generate_one:

        if st.button(
            "Generate Excel",
            type="primary",
            use_container_width=True,
        ):

            with st.spinner(f"Membuat file untuk {selected_employee}..."):

                try:

                    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

                    output_file = create_employee_file(
                        employee_name=selected_employee,
                        master_file=MASTER_FILE,
                        df_self=data["df_self"],
                        df_superior=data["df_superior"],
                        scored_answers=scored_answers_df,
                        indicator_master_df=data["indicator_master_df"],
                        output_dir=OUTPUT_DIR,
                    )

                    st.session_state.generated_file = output_file
                    st.session_state.generated_employee = selected_employee
                    st.session_state.generated_zip = None

                    st.success(
                        "Excel berhasil dibuat untuk "
                        f"{selected_employee}."
                    )

                except Exception as e:

                    st.session_state.generated_file = None

                    st.error(f"Gagal membuat Excel: {e}")

    # ========================================================
    # GENERATE ALL EMPLOYEES
    # ========================================================

    with col_generate_all:

        if st.button(
            "Generate All Employees",
            use_container_width=True,
        ):

            with st.spinner(f"Membuat {len(employees)} file employee..."):

                try:

                    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

                    zip_file, generated_files = create_all_employee_files(
                        employees=employees,
                        master_file=MASTER_FILE,
                        df_self=data["df_self"],
                        df_superior=data["df_superior"],
                        scored_answers=scored_answers_df,
                        indicator_master_df=data["indicator_master_df"],
                        output_dir=OUTPUT_DIR,
                    )

                    st.session_state.generated_zip = zip_file
                    st.session_state.generated_file = None
                    st.session_state.generated_employee = None

                    st.success(
                        f"{len(generated_files)} "
                        "file employee berhasil dibuat "
                        "dan dikumpulkan ke dalam ZIP."
                    )

                except Exception as e:

                    st.session_state.generated_zip = None

                    st.error(f"Gagal membuat seluruh file: {e}")


# ============================================================
# DOWNLOAD ONE EMPLOYEE
# ============================================================

if (
    st.session_state.generated_file is not None
    and st.session_state.generated_file.exists()
):

    output_file = st.session_state.generated_file

    render_html(
        """
        <div class="learning-result-card">
            <div class="learning-result-title">
                Download Employee Report
            </div>
            <div class="learning-result-description">
                File employee berhasil dibuat
                dan siap untuk di-download.
            </div>
        </div>
        """
    )

    with open(output_file, "rb") as f:
        file_bytes = f.read()

    st.download_button(
        label=f"Download {output_file.name}",
        data=file_bytes,
        file_name=output_file.name,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
    )


# ============================================================
# DOWNLOAD ALL EMPLOYEES ZIP
# ============================================================

if (
    st.session_state.generated_zip is not None
    and st.session_state.generated_zip.exists()
):

    zip_file = st.session_state.generated_zip

    render_html(
        """
        <div class="learning-result-card">
            <div class="learning-result-title">
                Download All Employees
            </div>
            <div class="learning-result-description">
                Seluruh employee report berhasil dibuat
                dan sudah dikumpulkan dalam satu ZIP.
            </div>
        </div>
        """
    )

    with open(zip_file, "rb") as f:
        zip_bytes = f.read()

    st.download_button(
        label="Download All Employees (.ZIP)",
        data=zip_bytes,
        file_name=zip_file.name,
        mime="application/zip",
        use_container_width=True,
    )