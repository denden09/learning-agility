# ============================================================
# api/processor.py
# LEARNING AGILITY PROCESSOR
# ============================================================

import io
import re
import zipfile
from pathlib import Path
from datetime import datetime, date

import numpy as np
import pandas as pd

from openpyxl import load_workbook


# ============================================================
# BASE DIRECTORY
# ============================================================

BASE_DIR = Path(
    r"C:\Users\zerox\OneDrive\Documents\Desktop\learning"
)


# ============================================================
# MASTER TEMPLATE
# ============================================================

MASTER_FILE = (
    BASE_DIR
    / "data"
    / "master"
    / "Learning agility self assesment & superior_24.7.2026.xlsx"
)


# ============================================================
# DIMENSIONS
# ============================================================

DIMENSIONS = {
    "mental_agility": "mental agility",
    "people_agility": "people agility",
    "change_agility": "change agility",
    "result_agility": "result agility",
    "self_awareness": "self awareness",
}


# ============================================================
# MASTER INPUT CELLS
# ============================================================

IDENTITY_INPUT_CELLS = [
    "D2",
    "D3",
    "D4",
    "D5",
    "D6",
    "D7",
    "D8",
    "D9",
    "D10",
    "D12",
]


ASSESSMENT_INPUT_COLUMNS = {
    "self_rating": 9,       # I
    "self_comment": 10,     # J
    "superior_rating": 11,  # K
    "superior_comment": 12, # L
}


ASSESSMENT_ROWS = list(
    range(13, 23)
)


ASSESSMENT_SHEETS = list(
    DIMENSIONS.values()
)


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(value):

    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    value = str(value)

    value = value.replace("\xa0", " ")
    value = value.replace("\n", " ")
    value = value.replace("\r", " ")

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip().lower()


# ============================================================
# CLEAN EXCEL VALUE
# ============================================================

def clean_excel_value(value):

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return None

    return value


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(
    df,
    candidates
):

    if df is None or df.empty:
        return None

    normalized_columns = {
        normalize_text(col): col
        for col in df.columns
    }

    for candidate in candidates:

        key = normalize_text(
            candidate
        )

        if key in normalized_columns:

            return normalized_columns[key]

    return None


# ============================================================
# WRITE VALUE ONLY
# ============================================================

def write_value_only(
    cell,
    value
):

    """
    Menulis VALUE saja.

    Tidak menyentuh:
    - font
    - fill
    - border
    - alignment
    - number format
    - protection
    """

    value = clean_excel_value(value)

    if value is None:

        cell.value = None

    else:

        cell.value = value


# ============================================================
# LOAD MASTER STRUCTURE
# ============================================================

def extract_indicator_master(
    master_wb
):

    indicator_master = []

    for dimension_key, sheet_name in DIMENSIONS.items():

        ws = master_wb[sheet_name]

        for indicator_no, row in enumerate(
            range(13, 23),
            start=1
        ):

            indicator = ws.cell(
                row,
                2
            ).value

            question = ws.cell(
                row,
                3
            ).value

            levels = {
                1: ws.cell(row, 4).value,
                2: ws.cell(row, 5).value,
                3: ws.cell(row, 6).value,
                4: ws.cell(row, 7).value,
                5: ws.cell(row, 8).value,
            }

            indicator_master.append({

                "dimension":
                    dimension_key,

                "sheet_name":
                    sheet_name,

                "indicator_no":
                    indicator_no,

                "indicator_id":
                    f"{dimension_key}_{indicator_no:02d}",

                "indicator":
                    indicator,

                "question":
                    question,

                "question_normalized":
                    normalize_text(question),

                "level_1":
                    levels[1],

                "level_2":
                    levels[2],

                "level_3":
                    levels[3],

                "level_4":
                    levels[4],

                "level_5":
                    levels[5],

                "level_1_normalized":
                    normalize_text(levels[1]),

                "level_2_normalized":
                    normalize_text(levels[2]),

                "level_3_normalized":
                    normalize_text(levels[3]),

                "level_4_normalized":
                    normalize_text(levels[4]),

                "level_5_normalized":
                    normalize_text(levels[5]),
            })

    return pd.DataFrame(
        indicator_master
    )


# ============================================================
# ANSWER MAPPING
# ============================================================

def build_answer_mapping(
    indicator_master_df
):

    answer_mapping = {}

    for _, row in indicator_master_df.iterrows():

        question_key = row[
            "question_normalized"
        ]

        answer_mapping[
            question_key
        ] = {

            normalize_text(
                row["level_1"]
            ): 1,

            normalize_text(
                row["level_2"]
            ): 2,

            normalize_text(
                row["level_3"]
            ): 3,

            normalize_text(
                row["level_4"]
            ): 4,

            normalize_text(
                row["level_5"]
            ): 5,
        }

    return answer_mapping


# ============================================================
# DETECT QUESTION COLUMNS
# ============================================================

def detect_question_columns(
    df,
    master_questions
):

    master_question_set = {
        normalize_text(q)
        for q in master_questions
        if pd.notna(q)
    }

    question_columns = []

    for col in df.columns:

        normalized_col = normalize_text(
            col
        )

        if normalized_col in master_question_set:

            question_columns.append(col)

    return question_columns


# ============================================================
# PROCESS ASSESSMENT
# ============================================================

def process_assessment(
    df,
    question_columns,
    rater_type,
    indicator_master_df,
    answer_mapping
):

    results = []

    for row_index, row in df.iterrows():

        row_dict = {
            normalize_text(k): v
            for k, v in row.items()
        }

        # ----------------------------------------------------
        # EMPLOYEE
        # ----------------------------------------------------

        if rater_type == "Self":

            employee_name = clean_excel_value(
                row.get(
                    "Nama",
                    row_dict.get(
                        "nama",
                        ""
                    )
                )
            )

            superior_name = ""

        else:

            employee_name = clean_excel_value(
                row.get(
                    "Nama bawahan yang dinilai",
                    row_dict.get(
                        "nama bawahan yang dinilai",
                        ""
                    )
                )
            )

            superior_name = clean_excel_value(
                row.get(
                    "Nama",
                    row_dict.get(
                        "nama",
                        ""
                    )
                )
            )

        # ----------------------------------------------------
        # IDENTITY
        # ----------------------------------------------------

        email = clean_excel_value(
            row.get(
                "Email",
                row_dict.get(
                    "email",
                    ""
                )
            )
        )

        division = clean_excel_value(
            row.get(
                "Divisi",
                row_dict.get(
                    "divisi",
                    ""
                )
            )
        )

        unit = clean_excel_value(
            row.get(
                "Unit",
                row_dict.get(
                    "unit",
                    ""
                )
            )
        )

        department = clean_excel_value(
            row.get(
                "Departemen",
                row_dict.get(
                    "departemen",
                    ""
                )
            )
        )

        directorate = clean_excel_value(
            row.get(
                "Direktorat",
                row_dict.get(
                    "direktorat",
                    ""
                )
            )
        )

        current_level = clean_excel_value(
            row.get(
                "Level Saat Ini",
                row_dict.get(
                    "level saat ini",
                    ""
                )
            )
        )

        target_level = clean_excel_value(
            row.get(
                "Level Tujuan",
                row_dict.get(
                    "level tujuan",
                    ""
                )
            )
        )

        next_path = clean_excel_value(
            row.get(
                "Level Next Path",
                row_dict.get(
                    "level next path",
                    ""
                )
            )
        )

        completion_time = row.get(
            "Completion time",
            row_dict.get(
                "completion time",
                None
            )
        )

        last_modified_time = row.get(
            "Last modified time",
            row_dict.get(
                "last modified time",
                None
            )
        )

        # ----------------------------------------------------
        # QUESTIONS
        # ----------------------------------------------------

        for question_col in question_columns:

            question_normalized = normalize_text(
                question_col
            )

            answer = row.get(
                question_col,
                ""
            )

            if pd.isna(answer):
                continue

            answer_str = str(
                answer
            ).strip()

            if not answer_str:
                continue

            answer_normalized = normalize_text(
                answer_str
            )

            master_rows = indicator_master_df[
                indicator_master_df[
                    "question_normalized"
                ]
                ==
                question_normalized
            ]

            # ------------------------------------------------
            # QUESTION NOT FOUND
            # ------------------------------------------------

            if master_rows.empty:

                results.append({

                    "rater_type":
                        rater_type,

                    "employee_name":
                        employee_name,

                    "superior_name":
                        superior_name,

                    "email":
                        email,

                    "division":
                        division,

                    "unit":
                        unit,

                    "department":
                        department,

                    "directorate":
                        directorate,

                    "current_level":
                        current_level,

                    "target_level":
                        target_level,

                    "next_path":
                        next_path,

                    "completion_time":
                        completion_time,

                    "last_modified_time":
                        last_modified_time,

                    "question":
                        question_col,

                    "answer_text":
                        answer_str,

                    "rating":
                        np.nan,

                    "mapping_status":
                        "QUESTION_NOT_FOUND",

                    "indicator_id":
                        "",

                    "dimension":
                        "",

                    "indicator_no":
                        np.nan,

                    "indicator":
                        "",
                })

                continue

            # ------------------------------------------------
            # MASTER QUESTION
            # ------------------------------------------------

            master_row = master_rows.iloc[0]

            rating = answer_mapping[
                question_normalized
            ].get(
                answer_normalized,
                np.nan
            )

            mapping_status = (
                "ANSWER_NOT_FOUND"
                if pd.isna(rating)
                else "MATCHED"
            )

            results.append({

                "rater_type":
                    rater_type,

                "employee_name":
                    employee_name,

                "superior_name":
                    superior_name,

                "email":
                    email,

                "division":
                    division,

                "unit":
                    unit,

                "department":
                    department,

                "directorate":
                    directorate,

                "current_level":
                    current_level,

                "target_level":
                    target_level,

                "next_path":
                    next_path,

                "completion_time":
                    completion_time,

                "last_modified_time":
                    last_modified_time,

                "question":
                    question_col,

                "answer_text":
                    answer_str,

                "rating":
                    rating,

                "mapping_status":
                    mapping_status,

                "indicator_id":
                    master_row[
                        "indicator_id"
                    ],

                "dimension":
                    master_row[
                        "dimension"
                    ],

                "indicator_no":
                    master_row[
                        "indicator_no"
                    ],

                "indicator":
                    master_row[
                        "indicator"
                    ],
            })

    return pd.DataFrame(
        results
    )


# ============================================================
# BUILD SCORING DATA
# ============================================================

def build_scoring(
    df_self,
    df_superior
):

    if not MASTER_FILE.exists():

        raise FileNotFoundError(
            f"Master file tidak ditemukan:\n{MASTER_FILE}"
        )

    # --------------------------------------------------------
    # LOAD MASTER
    # --------------------------------------------------------

    master_wb = load_workbook(
        MASTER_FILE,
        data_only=False
    )

    # --------------------------------------------------------
    # INDICATOR MASTER
    # --------------------------------------------------------

    indicator_master_df = extract_indicator_master(
        master_wb
    )

    if len(indicator_master_df) != 50:

        raise ValueError(
            "Master harus memiliki 50 indicator."
        )

    # --------------------------------------------------------
    # ANSWER MAPPING
    # --------------------------------------------------------

    answer_mapping = build_answer_mapping(
        indicator_master_df
    )

    # --------------------------------------------------------
    # QUESTION COLUMNS
    # --------------------------------------------------------

    master_questions = (
        indicator_master_df[
            "question"
        ]
        .tolist()
    )

    self_question_cols = (
        detect_question_columns(
            df_self,
            master_questions
        )
    )

    superior_question_cols = (
        detect_question_columns(
            df_superior,
            master_questions
        )
    )

    if len(self_question_cols) != 50:

        raise ValueError(
            "Self Assessment tidak memiliki "
            f"50 question yang cocok dengan master. "
            f"Ditemukan: {len(self_question_cols)}"
        )

    if len(superior_question_cols) != 50:

        raise ValueError(
            "Superior Assessment tidak memiliki "
            f"50 question yang cocok dengan master. "
            f"Ditemukan: {len(superior_question_cols)}"
        )

    # --------------------------------------------------------
    # PROCESS SELF
    # --------------------------------------------------------

    self_scored_df = process_assessment(
        df=df_self,
        question_columns=self_question_cols,
        rater_type="Self",
        indicator_master_df=indicator_master_df,
        answer_mapping=answer_mapping
    )

    # --------------------------------------------------------
    # PROCESS SUPERIOR
    # --------------------------------------------------------

    superior_scored_df = process_assessment(
        df=df_superior,
        question_columns=superior_question_cols,
        rater_type="Superior",
        indicator_master_df=indicator_master_df,
        answer_mapping=answer_mapping
    )

    # --------------------------------------------------------
    # COMBINE
    # --------------------------------------------------------

    scored_answers_df = pd.concat(
        [
            self_scored_df,
            superior_scored_df
        ],
        ignore_index=True
    )

    # --------------------------------------------------------
    # VALIDATE
    # --------------------------------------------------------

    if scored_answers_df.empty:

        raise ValueError(
            "Tidak ada data assessment yang berhasil diproses."
        )

    empty_employee = (
        scored_answers_df[
            "employee_name"
        ]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("")
        .sum()
    )

    if empty_employee > 0:

        raise ValueError(
            f"Terdapat {empty_employee} record "
            "dengan employee_name kosong."
        )

    # --------------------------------------------------------
    # INDICATOR SCORE
    # --------------------------------------------------------

    indicator_score_df = (
        scored_answers_df[
            scored_answers_df[
                "mapping_status"
            ]
            ==
            "MATCHED"
        ]
        .groupby(
            [
                "rater_type",
                "employee_name",
                "indicator_id"
            ],
            as_index=False
        )[
            "rating"
        ]
        .mean()
        .rename(
            columns={
                "rating":
                    "indicator_score"
            }
        )
    )

    indicator_score_df = (
        indicator_score_df.merge(
            indicator_master_df[
                [
                    "indicator_id",
                    "dimension",
                    "indicator_no",
                    "indicator"
                ]
            ],
            on="indicator_id",
            how="left"
        )
    )

    # --------------------------------------------------------
    # DIMENSION SCORE
    # --------------------------------------------------------

    dimension_score_df = (
        indicator_score_df
        .groupby(
            [
                "rater_type",
                "employee_name",
                "dimension"
            ],
            as_index=False
        )[
            "indicator_score"
        ]
        .mean()
        .rename(
            columns={
                "indicator_score":
                    "dimension_score"
            }
        )
    )

    # --------------------------------------------------------
    # GAP
    # --------------------------------------------------------

    dimension_gap_df = (
        dimension_score_df
        .pivot_table(
            index=[
                "employee_name",
                "dimension"
            ],
            columns="rater_type",
            values="dimension_score"
        )
        .reset_index()
    )

    if "Self" not in dimension_gap_df.columns:

        dimension_gap_df["Self"] = np.nan

    if "Superior" not in dimension_gap_df.columns:

        dimension_gap_df["Superior"] = np.nan

    dimension_gap_df["gap"] = (
        dimension_gap_df["Self"]
        -
        dimension_gap_df["Superior"]
    )

    return {
        "indicator_master_df":
            indicator_master_df,

        "scored_answers_df":
            scored_answers_df,

        "indicator_score_df":
            indicator_score_df,

        "dimension_score_df":
            dimension_score_df,

        "dimension_gap_df":
            dimension_gap_df,
    }


# ============================================================
# FILL MASTER FOR ONE EMPLOYEE
# ============================================================

def fill_master_for_employee(
    employee_name,
    master_file,
    scored_answers,
    df_self,
    df_superior
):

    # ========================================================
    # LOAD MASTER
    # ========================================================

    wb = load_workbook(
        master_file,
        data_only=False
    )

    # ========================================================
    # TARGET EMPLOYEE
    # ========================================================

    target_employee = normalize_text(
        employee_name
    )

    # ========================================================
    # FILTER EMPLOYEE DATA
    # ========================================================

    employee_data = scored_answers[
        scored_answers[
            "employee_name"
        ]
        .apply(normalize_text)
        ==
        target_employee
    ].copy()

    if employee_data.empty:

        raise ValueError(
            f"Tidak ada data assessment "
            f"untuk employee: {employee_name}"
        )

    # ========================================================
    # IDENTITAS
    # ========================================================

    ws_identity = wb[
        "Identitas"
    ]

    # --------------------------------------------------------
    # D2 — NAMA
    # --------------------------------------------------------

    write_value_only(
        ws_identity["D2"],
        employee_name
    )

    # ========================================================
    # CARI DATA SELF RAW
    # ========================================================

    self_raw = pd.DataFrame()

    if (
        df_self is not None
        and not df_self.empty
    ):

        self_name_col = find_column(
            df_self,
            [
                "Nama",
                "Name"
            ]
        )

        if self_name_col is not None:

            self_raw = df_self[
                df_self[
                    self_name_col
                ]
                .apply(normalize_text)
                ==
                target_employee
            ].copy()

    # ========================================================
    # CARI NAMA SUPERIOR
    # ========================================================

    superior_name = None

    if (
        df_superior is not None
        and not df_superior.empty
    ):

        subordinate_col = find_column(
            df_superior,
            [
                "Nama bawahan yang dinilai",
                "Nama Bawahan",
                "Bawahan"
            ]
        )

        if subordinate_col is not None:

            superior_employee_data = (
                df_superior[
                    df_superior[
                        subordinate_col
                    ]
                    .apply(normalize_text)
                    ==
                    target_employee
                ]
                .copy()
            )

            if not superior_employee_data.empty:

                superior_name_col = find_column(
                    df_superior,
                    [
                        "Nama",
                        "Name",
                        "Nama atasan yang menilai",
                        "Nama atasan",
                        "Nama superior",
                        "Superior Name"
                    ]
                )

                if superior_name_col is not None:

                    superior_name = clean_excel_value(
                        superior_employee_data.iloc[
                            0
                        ][
                            superior_name_col
                        ]
                    )

    # --------------------------------------------------------
    # D3
    # --------------------------------------------------------

    write_value_only(
        ws_identity["D3"],
        superior_name
    )

    # ========================================================
    # IDENTITY MAPPING
    # ========================================================

    identity_mapping = {

        "D4": [
            "Level Saat Ini",
            "Level saat ini",
            "Current Level"
        ],

        "D5": [
            "Level Tujuan",
            "Level tujuan",
            "Target Level"
        ],

        "D6": [
            "Tanggal Assesment",
            "Tanggal Assessment",
            "Assessment Date",
            "Completion time",
            "Start time",
            "Last modified time"
        ],

        "D7": [
            "Unit"
        ],

        "D8": [
            "Departemen",
            "Department"
        ],

        "D9": [
            "Divisi",
            "Division"
        ],

        "D10": [
            "Direktorat",
            "Directorate"
        ],

        "D12": [
            "Level Next Path",
            "Next Path",
            "Level Next"
        ]
    }

    # ========================================================
    # ISI IDENTITY DARI SELF RAW
    # ========================================================

    if not self_raw.empty:

        self_row = self_raw.iloc[0]

        for target_cell, candidates in identity_mapping.items():

            source_col = find_column(
                self_raw,
                candidates
            )

            if source_col is None:
                continue

            value = clean_excel_value(
                self_row[source_col]
            )

            if value is None:
                continue

            if target_cell == "D6":

                if isinstance(
                    value,
                    pd.Timestamp
                ):

                    value = value.to_pydatetime()

            write_value_only(
                ws_identity[target_cell],
                value
            )

    # ========================================================
    # PROCESS 5 DIMENSIONS
    # ========================================================

    for dimension_key, sheet_name in DIMENSIONS.items():

        ws = wb[
            sheet_name
        ]

        dimension_data = employee_data[
            employee_data[
                "dimension"
            ]
            ==
            dimension_key
        ].copy()

        # ----------------------------------------------------
        # 10 INDICATORS
        # ----------------------------------------------------

        for indicator_no, row in enumerate(
            ASSESSMENT_ROWS,
            start=1
        ):

            indicator_id = (
                f"{dimension_key}_{indicator_no:02d}"
            )

            indicator_data = dimension_data[
                dimension_data[
                    "indicator_id"
                ]
                ==
                indicator_id
            ].copy()

            # =================================================
            # SELF
            # =================================================

            self_data = indicator_data[
                indicator_data[
                    "rater_type"
                ]
                ==
                "Self"
            ].copy()

            self_rating = None

            if not self_data.empty:

                ratings = pd.to_numeric(
                    self_data[
                        "rating"
                    ],
                    errors="coerce"
                ).dropna()

                if not ratings.empty:

                    self_rating = ratings.iloc[0]

            # ------------------------------------------------
            # I — SELF RATING
            # ------------------------------------------------

            write_value_only(
                ws.cell(
                    row=row,
                    column=ASSESSMENT_INPUT_COLUMNS[
                        "self_rating"
                    ]
                ),
                self_rating
            )

            # ------------------------------------------------
            # J — SELF ANSWER
            # ------------------------------------------------

            self_comment = None

            if not self_data.empty:

                self_answers = (
                    self_data[
                        "answer_text"
                    ]
                    .dropna()
                    .astype(str)
                    .str.strip()
                )

                self_answers = self_answers[
                    self_answers != ""
                ]

                if not self_answers.empty:

                    self_comment = "\n".join(
                        self_answers.tolist()
                    )

            write_value_only(
                ws.cell(
                    row=row,
                    column=ASSESSMENT_INPUT_COLUMNS[
                        "self_comment"
                    ]
                ),
                self_comment
            )

            # =================================================
            # SUPERIOR
            # =================================================

            superior_data = indicator_data[
                indicator_data[
                    "rater_type"
                ]
                ==
                "Superior"
            ].copy()

            superior_rating = None

            if not superior_data.empty:

                ratings = pd.to_numeric(
                    superior_data[
                        "rating"
                    ],
                    errors="coerce"
                ).dropna()

                if not ratings.empty:

                    superior_rating = ratings.iloc[0]

            # ------------------------------------------------
            # K — SUPERIOR RATING
            # ------------------------------------------------

            write_value_only(
                ws.cell(
                    row=row,
                    column=ASSESSMENT_INPUT_COLUMNS[
                        "superior_rating"
                    ]
                ),
                superior_rating
            )

            # ------------------------------------------------
            # L — SUPERIOR ANSWER
            # ------------------------------------------------

            superior_comment = None

            if not superior_data.empty:

                superior_answers = (
                    superior_data[
                        "answer_text"
                    ]
                    .dropna()
                    .astype(str)
                    .str.strip()
                )

                superior_answers = superior_answers[
                    superior_answers != ""
                ]

                if not superior_answers.empty:

                    superior_comment = "\n".join(
                        superior_answers.tolist()
                    )

            write_value_only(
                ws.cell(
                    row=row,
                    column=ASSESSMENT_INPUT_COLUMNS[
                        "superior_comment"
                    ]
                ),
                superior_comment
            )

            # =================================================
            # PENTING
            # =================================================
            #
            # TIDAK MENYENTUH:
            #
            # M13:O22
            # N13:N22
            # O13:O22
            #
            # termasuk ROW 23.
            #
            # Formula dan formatting tetap dari MASTER.
            # =================================================

    return wb


# ============================================================
# PERFORMANCE LEVEL
# ============================================================

def performance_level_from_score(
    score
):

    if score is None:
        return ""

    try:

        score = float(score)

    except (
        ValueError,
        TypeError
    ):

        return ""

    if score >= 4.5:

        return "Excellent"

    elif score >= 4.0:

        return "Strong"

    elif score >= 3.0:

        return "Meet"

    else:

        return "Needs Development"


# ============================================================
# PROFILE SUMMARY — GAP
# ============================================================

def fill_profile_summary_gap(
    wb
):

    if "Profile Summary" not in wb.sheetnames:

        return wb

    ws = wb[
        "Profile Summary"
    ]

    dimension_rows = {
        "mental_agility": 98,
        "people_agility": 99,
        "change_agility": 100,
        "result_agility": 101,
        "self_awareness": 102,
    }

    for dimension_key, row in dimension_rows.items():

        # ----------------------------------------------------
        # B = SELF
        # C = SUPERIOR
        # D = GAP
        # ----------------------------------------------------

        self_score = None
        superior_score = None

        try:

            self_score = ws.cell(
                row=row,
                column=2
            ).value

            superior_score = ws.cell(
                row=row,
                column=3
            ).value

        except Exception:

            continue

        if (
            self_score is not None
            and superior_score is not None
        ):

            try:

                gap = (
                    float(self_score)
                    -
                    float(superior_score)
                )

                ws.cell(
                    row=row,
                    column=4
                ).value = gap

            except (
                ValueError,
                TypeError
            ):

                ws.cell(
                    row=row,
                    column=4
                ).value = None

        else:

            ws.cell(
                row=row,
                column=4
            ).value = None

    return wb


# ============================================================
# PROFILE SUMMARY — OVERALL
# ============================================================

def fill_profile_summary_summary(
    wb
):

    if "Profile Summary" not in wb.sheetnames:

        return wb

    ws = wb[
        "Profile Summary"
    ]

    self_values = []
    superior_values = []

    # --------------------------------------------------------
    # FIVE DIMENSIONS
    # --------------------------------------------------------

    for row in range(
        98,
        103
    ):

        self_score = ws.cell(
            row=row,
            column=2
        ).value

        superior_score = ws.cell(
            row=row,
            column=3
        ).value

        if self_score is not None:

            try:

                self_values.append(
                    float(self_score)
                )

            except (
                ValueError,
                TypeError
            ):

                pass

        if superior_score is not None:

            try:

                superior_values.append(
                    float(superior_score)
                )

            except (
                ValueError,
                TypeError
            ):

                pass

    # --------------------------------------------------------
    # OVERALL SELF
    # --------------------------------------------------------

    if self_values:

        overall_self = np.mean(
            self_values
        )

        ws["B103"] = float(
            overall_self
        )

    else:

        ws["B103"] = None

    # --------------------------------------------------------
    # OVERALL SUPERIOR
    # --------------------------------------------------------

    if superior_values:

        overall_superior = np.mean(
            superior_values
        )

        ws["C103"] = float(
            overall_superior
        )

    else:

        ws["C103"] = None

    # --------------------------------------------------------
    # OVERALL GAP
    # --------------------------------------------------------

    if (
        ws["B103"].value is not None
        and
        ws["C103"].value is not None
    ):

        overall_gap = (
            float(ws["B103"].value)
            -
            float(ws["C103"].value)
        )

        ws["D103"] = overall_gap

        ws["E103"] = classify_dimension_gap(
            overall_gap
        )

    else:

        ws["D103"] = None
        ws["E103"] = None

    return wb


# ============================================================
# GAP CLASSIFICATION
# ============================================================

def classify_dimension_gap(
    gap
):

    if gap is None:
        return ""

    try:

        gap = float(gap)

    except (
        ValueError,
        TypeError
    ):

        return ""

    if gap >= 0.50:

        return "Self Higher"

    elif gap <= -0.50:

        return "Superior Higher"

    else:

        return "Aligned"


# ============================================================
# RADAR CHART DATA
# ============================================================

def fill_radar_chart_data(
    wb
):

    if "Profile Summary" not in wb.sheetnames:

        return wb

    ws = wb[
        "Profile Summary"
    ]

    radar_config = [

        (
            "Mental Agility",
            98,
            109
        ),

        (
            "People Agility",
            99,
            110
        ),

        (
            "Change Agility",
            100,
            111
        ),

        (
            "Result Agility",
            101,
            112
        ),

        (
            "Self Awareness",
            102,
            113
        ),
    ]

    for (
        dimension_name,
        source_row,
        target_row
    ) in radar_config:

        # ----------------------------------------------------
        # DIMENSION
        # ----------------------------------------------------

        ws.cell(
            row=target_row,
            column=1
        ).value = dimension_name

        # ----------------------------------------------------
        # SCORE
        # Superior score
        # ----------------------------------------------------

        ws.cell(
            row=target_row,
            column=2
        ).value = ws.cell(
            row=source_row,
            column=3
        ).value

        # ----------------------------------------------------
        # PERFORMANCE LEVEL
        # ----------------------------------------------------

        score = ws.cell(
            row=target_row,
            column=2
        ).value

        ws.cell(
            row=target_row,
            column=3
        ).value = (
            performance_level_from_score(
                score
            )
        )

    return wb


# ============================================================
# BUILD PROFILE SUMMARY
# ============================================================

def build_profile_summary(
    wb,
    employee_name
):

    wb = fill_profile_summary_gap(
        wb
    )

    wb = fill_profile_summary_summary(
        wb
    )

    wb = fill_radar_chart_data(
        wb
    )

    return wb


# ============================================================
# WORKBOOK → BYTES
# ============================================================

def workbook_to_bytes(
    wb
):

    output = io.BytesIO()

    wb.save(
        output
    )

    output.seek(0)

    return output.getvalue()


# ============================================================
# SAFE FILE NAME
# ============================================================

def safe_filename(
    employee_name
):

    filename = re.sub(
        r'[\\/*?:"<>|]',
        "_",
        str(employee_name)
    )

    filename = filename.strip()

    if not filename:

        filename = "Employee"

    return filename


# ============================================================
# GENERATE REPORTS
# ============================================================

def generate_reports(
    df_self,
    df_superior
):

    # ========================================================
    # VALIDATE MASTER
    # ========================================================

    if not MASTER_FILE.exists():

        raise FileNotFoundError(
            f"Master file tidak ditemukan:\n{MASTER_FILE}"
        )

    # ========================================================
    # BUILD SCORING
    # ========================================================

    scoring = build_scoring(
        df_self=df_self,
        df_superior=df_superior
    )

    scored_answers_df = scoring[
        "scored_answers_df"
    ]

    # ========================================================
    # EMPLOYEE LIST
    # ========================================================

    employees = sorted(

        scored_answers_df[
            scored_answers_df[
                "rater_type"
            ]
            ==
            "Self"
        ][
            "employee_name"
        ]

        .dropna()

        .astype(str)

        .str.strip()

        .loc[
            lambda s: s != ""
        ]

        .unique()
    )

    if len(employees) == 0:

        raise ValueError(
            "Tidak ada employee Self Assessment."
        )

    # ========================================================
    # GENERATE
    # ========================================================

    generated = []
    failed = []

    for employee_name in employees:

        try:

            # ------------------------------------------------
            # CREATE MASTER
            # ------------------------------------------------

            wb = fill_master_for_employee(

                employee_name=employee_name,

                master_file=MASTER_FILE,

                scored_answers=scored_answers_df,

                df_self=df_self,

                df_superior=df_superior
            )

            # ------------------------------------------------
            # PROFILE SUMMARY
            # ------------------------------------------------

            wb = build_profile_summary(
                wb,
                employee_name
            )

            # ------------------------------------------------
            # SAVE TO MEMORY
            # ------------------------------------------------

            output_bytes = workbook_to_bytes(
                wb
            )

            wb.close()

            filename = (
                "Learning_Agility_"
                +
                safe_filename(
                    employee_name
                )
                +
                ".xlsx"
            )

            generated.append({

                "employee":
                    employee_name,

                "filename":
                    filename,

                "content":
                    output_bytes,
            })

        except Exception as exc:

            failed.append({

                "employee":
                    employee_name,

                "error":
                    repr(exc),
            })

    return {
        "generated":
            generated,

        "failed":
            failed,

        "employee_count":
            len(employees),

        "master_file":
            str(MASTER_FILE),
    }


# ============================================================
# CREATE ZIP
# ============================================================

def create_reports_zip(
    generated_files
):

    if not generated_files:

        raise ValueError(
            "Tidak ada report yang berhasil dibuat."
        )

    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(
        zip_buffer,
        mode="w",
        compression=zipfile.ZIP_DEFLATED
    ) as zf:

        for item in generated_files:

            zf.writestr(
                item["filename"],
                item["content"]
            )

    zip_buffer.seek(0)

    return zip_buffer.getvalue()